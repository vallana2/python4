"""Excel / CSV export and product-picture downloads."""
import csv
import io
from datetime import datetime

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config

# Required columns (values are in Rwandan francs) + extra USD reference columns
COLUMNS = [
    ("Product_Name", "product_name"),
    ("Product_Price", "product_price"),
    ("Commission_fee", "commission_fee"),
    ("Total_price", "total_price"),
    ("global_rating", "global_rating"),
    ("retailer_name", "retailer_name"),
    ("Retailer_email", "retailer_email"),
    ("Product_picture_name", "product_picture_name"),
    ("Currency", "currency"),
    ("Price_USD", "price_usd"),
    ("Total_price_USD", "total_usd"),
]
MONEY_FIELDS = {"product_price", "commission_fee", "total_price"}


def download_pictures(products):
    """Save each product picture in Online_Pictures and set product_picture_name."""
    config.PICTURES_DIR.mkdir(exist_ok=True)
    for index, product in enumerate(products, start=1):
        key = product.get("asin") or f"item{index}"
        filename = f"{key}.jpg"
        path = config.PICTURES_DIR / filename
        product["product_picture_name"] = filename

        if path.exists():
            continue
        if not product.get("image_url"):
            product["product_picture_name"] = "Not available"
            continue
        try:
            response = requests.get(product["image_url"], timeout=10)
            response.raise_for_status()
            path.write_bytes(response.content)
        except requests.exceptions.RequestException:
            product["product_picture_name"] = "Not available"   # keep going
    return products


def _row(product):
    values = []
    for _, field in COLUMNS:
        value = "RWF" if field == "currency" else product.get(field)
        values.append("N/A" if value is None else value)
    return values


def build_workbook(products, term="", rate=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append([header for header, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B6E4F")
        cell.alignment = Alignment(horizontal="center")

    for product in products:
        ws.append(_row(product))
    for row in ws.iter_rows(min_row=2):
        for cell, (_, field) in zip(row, COLUMNS):
            if field in MONEY_FIELDS:
                cell.number_format = '#,##0.00'

    for col_index in range(1, len(COLUMNS) + 1):
        letter = get_column_letter(col_index)
        widest = max(len(str(c.value)) for c in ws[letter])
        ws.column_dimensions[letter].width = min(widest + 3, 60)
    ws.freeze_panes = "A2"

    summary = wb.create_sheet("Summary")
    summary.append(["Search term", term])
    summary.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")])
    summary.append(["Exchange rate (1 USD)", f"{rate:,.2f} RWF" if rate else "N/A"])
    summary.append(["Products", len(products)])
    summary.append(["Total commission (RWF)", round(sum(p["commission_fee"] for p in products), 2)])
    if products and products[0].get("sample"):
        summary.append(["Data source", "SAMPLE DATA - Amazon returned no prices"])
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 28
    for row in summary.iter_rows():
        row[0].font = Font(bold=True)
    return wb


def save_excel(products, term="", rate=None):
    """Write the product list to Desktop/FirstName_LastName_RegNumber.xlsx."""
    path = config.get_desktop_path() / config.EXCEL_FILENAME
    wb = build_workbook(products, term, rate)
    try:
        wb.save(path)
    except PermissionError:
        raise PermissionError(
            f"Could not save {path.name} on the Desktop. Close it in Excel and search again.")
    return path


def excel_bytes(products, term="", rate=None):
    """Same workbook as an in-memory file, for the browser Download button."""
    buffer = io.BytesIO()
    build_workbook(products, term, rate).save(buffer)
    buffer.seek(0)
    return buffer


def csv_bytes(products):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([header for header, _ in COLUMNS])
    for product in products:
        writer.writerow(_row(product))
    return io.BytesIO(buffer.getvalue().encode("utf-8-sig"))     # opens correctly in Excel